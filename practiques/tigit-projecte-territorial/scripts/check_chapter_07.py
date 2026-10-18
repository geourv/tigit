#!/usr/bin/env python3
"""Check the chapter 7 palette workbook and vector proofs."""

import argparse
import re
from xml.etree import ElementTree

from openpyxl import load_workbook


from practice_paths import project_root
from check_chapter_03 import check_bar_axes

PROJECT_ROOT = project_root()
WORKBOOK = PROJECT_ROOT / "data" / "processed" / "tigit-07-teoria-color-teaching.xlsx"
FIGURES = PROJECT_ROOT / "outputs" / "figures" / "chapter-07"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--figures", action="store_true", help="Also require every declared vector proof")
    args = parser.parse_args()
    workbook = load_workbook(WORKBOOK, data_only=False)
    assert "palette" in workbook.sheetnames
    assert workbook["project"]["B3"].value == "07"
    assert workbook["palette"].max_row >= 26
    assert str(workbook["palette"]["F2"].value).startswith("=IF(")
    assert str(workbook["palette"]["L2"].value).startswith("=IF(")
    palette = workbook["palette"]
    roles = {palette.cell(row, 3).value: palette.cell(row, 5).value for row in range(2, palette.max_row + 1)}
    assert roles["young"] == "#EDF8B1"
    assert roles["working_age"] == "#7FCDBB"
    assert roles["older"] == "#2C7FB8"
    assert roles["female"] == "#CA0020"
    assert roles["male"] == "#0571B0"
    for row in range(2, palette.max_row + 1):
        hex_code = palette.cell(row, 5).value
        grayscale_check = palette.cell(row, 17).value
        cvd_check = palette.cell(row, 18).value
        notes = palette.cell(row, 19).value
        assert "PENDING" not in (grayscale_check, cvd_check)
        if hex_code:
            assert grayscale_check != "NO APLICA: falta codi HEX"
            assert cvd_check != "NO APLICA: falta codi HEX"
            assert len(re.findall(r"#[0-9A-F]{6}", notes or "")) == 4
        else:
            assert grayscale_check == "NO APLICA: falta codi HEX"
            assert cvd_check == "NO APLICA: falta codi HEX"
    for role in ("young", "working_age", "older", "class_1", "class_2", "class_3", "class_4", "class_5"):
        row = next(row for row in range(2, palette.max_row + 1) if palette.cell(row, 3).value == role)
        assert str(palette.cell(row, 17).value).startswith("PASSA:")
        assert str(palette.cell(row, 18).value).startswith("PASSA:")
    assert len(workbook["pivot_county_control"]._pivots) == 1
    assert len(workbook["pivot_population_age_sex"]._pivots) == 1
    check_bar_axes(workbook)
    if args.figures:
        check_figures()
    print("Chapter 7 checks passed" + (" (including figures)" if args.figures else " (workbook only)"))


def check_figures() -> None:
    for stem in (
        "palette-proof-age-structure-tarragones-2021",
        "palette-proof-non-principal-housing-tarragones-2021",
        "palette-proof-population-pyramid-vila-seca-2021",
    ):
        assert (FIGURES / f"{stem}.pdf").stat().st_size > 10_000
        ElementTree.parse(FIGURES / f"{stem}.svg")


if __name__ == "__main__":
    main()
