#!/usr/bin/env python3
"""Check chapter 3 workbook and vector exports."""

import argparse
from xml.etree import ElementTree

from openpyxl import load_workbook


from practice_paths import project_root

PROJECT_ROOT = project_root()
WORKBOOK = PROJECT_ROOT / "data" / "processed" / "tigit-03-semiologia-visualitzacio-teaching.xlsx"
FIGURES = PROJECT_ROOT / "outputs" / "figures" / "chapter-03"


def axis_title(axis) -> str:
    assert axis.title is not None, "Axis title is missing"
    return "".join(node.text or "" for node in axis.title.to_tree().iter() if node.tag.rsplit("}", 1)[-1] == "t")


def check_bar_axes(workbook) -> None:
    age = workbook["chart_01_age_structure"]._charts[0]
    assert axis_title(age.x_axis) == "Municipi"
    assert axis_title(age.y_axis) == "% de població"
    for chart in [*workbook["chart_00_audit"]._charts, workbook["chart_02_nonprincipal"]._charts[0]]:
        assert axis_title(chart.x_axis) == "Municipi"
        assert axis_title(chart.y_axis) == "% sobre el total d'habitatges"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--figures", action="store_true", help="Also require every declared vector export")
    args = parser.parse_args()
    workbook = load_workbook(WORKBOOK, data_only=False)
    expected = {
        "charts_data", "chart_00_audit", "chart_01_age_structure",
        "chart_02_nonprincipal", "chart_03_scatter", "chart_04_housing_donut",
        "chart_05_population_time", "source_population_time", "prepared_population_time",
        "chart_06_population_pyramid", "source_population_pyramid", "prepared_population_pyramid",
        "pivot_population_age_sex", "pyramid_data", "chart_07_housing_histogram", "histogram_data",
    }
    assert expected.issubset(workbook.sheetnames)
    assert workbook["project"]["B3"].value == "03"
    assert workbook["charts_data"].max_row == 23
    assert len(workbook["chart_00_audit"]._charts) == 2
    assert len(workbook["chart_01_age_structure"]._charts) == 1
    assert len(workbook["chart_02_nonprincipal"]._charts) == 1
    assert len(workbook["chart_03_scatter"]._charts) == 1
    assert len(workbook["chart_04_housing_donut"]._charts) == 1
    assert len(workbook["chart_05_population_time"]._charts) == 1
    assert len(workbook["chart_06_population_pyramid"]._charts) == 1
    assert len(workbook["chart_07_housing_histogram"]._charts) == 1
    assert len(workbook["pivot_county_control"]._pivots) == 1
    assert len(workbook["pivot_population_age_sex"]._pivots) == 1
    assert workbook["prepared_population_time"].max_row == 24
    assert workbook["pyramid_data"].max_row == 19
    scatter_series = workbook["chart_03_scatter"]._charts[0].series[0]
    assert scatter_series.graphicalProperties.line.noFill is True
    check_bar_axes(workbook)

    if args.figures:
        check_figures()
    print("Chapter 3 checks passed" + (" (including figures)" if args.figures else " (workbook only)"))


def check_figures() -> None:

    for name in (
        "age-structure-tarragones-2021.pdf",
        "non-principal-housing-tarragones-2021.pdf",
        "ageing-vs-non-principal-housing-tarragones-2021.pdf",
        "housing-composition-tarragones-2021.pdf",
        "population-vila-seca-2000-2022.pdf",
        "population-pyramid-vila-seca-2021.pdf",
        "housing-non-main-histogram-tarragones-2021.pdf",
    ):
        path = FIGURES / name
        assert path.stat().st_size > 10_000
    for name in (
        "age-structure-tarragones-2021.svg",
        "non-principal-housing-tarragones-2021.svg",
        "population-vila-seca-2000-2022.svg",
        "population-pyramid-vila-seca-2021.svg",
    ):
        ElementTree.parse(FIGURES / name)


if __name__ == "__main__":
    main()
