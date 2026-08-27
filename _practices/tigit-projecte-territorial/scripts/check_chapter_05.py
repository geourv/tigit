#!/usr/bin/env python3
"""Check the chapter 5 workbook and QGIS transfer table."""

import csv
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = PROJECT_ROOT / "data" / "processed" / "tigit-05-integracio-sig-teaching.xlsx"
CSV_OUTPUT = PROJECT_ROOT / "data" / "processed" / "municipal-indicators-tarragones-2021.csv"
CSVT_OUTPUT = CSV_OUTPUT.with_suffix(".csvt")
EXPECTED_HEADERS = [
    "mun_code", "municipality", "county_code", "year", "population_total",
    "population_65_plus", "population_65_plus_pct", "housing_total",
    "housing_non_main", "housing_non_main_pct", "indicator_status",
]


def main() -> None:
    workbook = load_workbook(WORKBOOK, data_only=False)
    assert workbook["project"]["B3"].value == "05"
    map_export = workbook["map_export"]
    assert map_export.max_row == 23
    assert [cell.value for cell in map_export[1]] == EXPECTED_HEADERS
    assert str(map_export["A2"].value).startswith("=municipal!")
    assert str(map_export["G2"].value).startswith("=indicators_demography!")
    assert str(map_export["J2"].value).startswith("=indicators_housing!")
    assert str(map_export["K2"].value).startswith("=IF(")

    with CSV_OUTPUT.open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 22
    assert set(rows[0]) == set(EXPECTED_HEADERS)
    assert len({row["mun_code"] for row in rows}) == 22
    assert any(row["mun_code"] == "431711" and row["municipality"] == "Vila-seca" for row in rows)
    assert {row["indicator_status"] for row in rows} == {"ok"}
    assert CSVT_OUTPUT.read_text(encoding="ascii").count(",") == len(EXPECTED_HEADERS) - 1
    print("Chapter 5 checks passed")


if __name__ == "__main__":
    main()
